"""TextBoxBuilder: precise text rendering with OXML paragraph/run control."""
from lxml import etree
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from ..schema.content import TextBlock, TextRun
from ..schema.design import DesignTokens


_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

def _a(tag: str) -> str:
    return f"{{{_A_NS}}}{tag}"


ALIGN_MAP = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
}


def apply_text_block(tf, block: TextBlock, tokens: DesignTokens) -> None:
    """Render a TextBlock into a python-pptx text frame, replacing existing content."""
    tf.clear()
    tf.word_wrap = True

    para = tf.paragraphs[0]
    _apply_paragraph_format(para, block, tokens)

    first = True
    for run_spec in block.runs:
        if first:
            run = para.add_run()
            first = False
        else:
            # Check for newline in text — start new paragraph
            if "\n" in run_spec.text:
                parts = run_spec.text.split("\n")
                for i, part in enumerate(parts):
                    if i == 0:
                        run = para.add_run()
                        _apply_run_format(run, run_spec, tokens, override_text=part)
                    else:
                        para = tf.add_paragraph()
                        _apply_paragraph_format(para, block, tokens)
                        if part:
                            run = para.add_run()
                            _apply_run_format(run, run_spec, tokens, override_text=part)
                continue
            run = para.add_run()

        _apply_run_format(run, run_spec, tokens)


def _apply_paragraph_format(para, block: TextBlock, tokens: DesignTokens) -> None:
    pPr = para._pPr
    if pPr is None:
        pPr = para._p.get_or_add_pPr()

    # Line spacing
    lnSpc = etree.SubElement(pPr, _a("lnSpc"))
    spcPct = etree.SubElement(lnSpc, _a("spcPct"))
    spcPct.set("val", str(block.line_spacing_pct * 1000))

    # Space before
    spcBef = etree.SubElement(pPr, _a("spcBef"))
    spcPts = etree.SubElement(spcBef, _a("spcPts"))
    spcPts.set("val", str(block.space_before_pt * 100))

    para.alignment = ALIGN_MAP.get(block.alignment, PP_ALIGN.LEFT)


def _apply_run_format(run, spec: TextRun, tokens: DesignTokens,
                      override_text: str | None = None) -> None:
    run.text = override_text if override_text is not None else spec.text
    run.font.bold = spec.bold
    run.font.italic = spec.italic
    run.font.name = tokens.typography.body_font

    size = spec.size_pt or tokens.typography.body_size
    run.font.size = Pt(size)

    if spec.color_hex:
        r, g, b = bytes.fromhex(spec.color_hex.lstrip("#"))
        run.font.color.rgb = RGBColor(r, g, b)
    else:
        r, g, b = bytes.fromhex(tokens.colors.text_primary.lstrip("#"))
        run.font.color.rgb = RGBColor(r, g, b)
